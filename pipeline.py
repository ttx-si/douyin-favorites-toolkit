"""
抖音收藏夹提取与逐字稿归档流水线 (Douyin Favorites Toolkit)
"""

import os
import sys
import json
import re
import time
import argparse
import urllib.request
import subprocess
from datetime import datetime

import imageio_ffmpeg
import numpy as np
from faster_whisper import WhisperModel
from rapidocr_onnxruntime import RapidOCR
from playwright.sync_api import sync_playwright
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

sys.stdout.reconfigure(encoding='utf-8')
ffmpeg_bin = imageio_ffmpeg.get_ffmpeg_exe()


def format_count(n):
    if not n:
        return "0"
    if n >= 10000:
        return f"{round(n / 10000, 1)}万"
    return str(n)


def format_duration(sec):
    if not sec:
        return "00:00"
    m = int(sec // 60)
    s = int(sec % 60)
    return f"{m:02d}:{s:02d}"


def sanitize_filename(name):
    clean = re.sub(r'[\\/*?:"<>|#\s\r\n]+', '_', name).strip('_')
    return clean[:35] if clean else "未知"


def load_audio_via_ffmpeg(url):
    """
    通过 FFmpeg 内存流下载并解码流媒体音轨为 16kHz float32 numpy 数组
    """
    cmd = [
        ffmpeg_bin, "-y",
        "-headers", "User-Agent: Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36\r\nReferer: https://www.douyin.com/\r\n",
        "-i", url,
        "-vn",
        "-f", "s16le",
        "-ar", "16000",
        "-ac", "1",
        "-"
    ]
    process = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    stdout, stderr = process.communicate()
    if process.returncode != 0:
        raise RuntimeError(f"FFmpeg failed with code {process.returncode}: {stderr.decode(errors='ignore')[-200:]}")
    return np.frombuffer(stdout, np.int16).flatten().astype(np.float32) / 32768.0


def process_video_asr(model, play_url):
    audio_np = load_audio_via_ffmpeg(play_url)
    duration_s = len(audio_np) / 16000
    segments, info = model.transcribe(audio_np, language="zh", beam_size=1)

    transcript_blocks = []
    curr_text = ""
    curr_start = 0.0
    curr_end = 0.0

    for s in segments:
        text = s.text.strip()
        if not text:
            continue
        if not curr_text:
            curr_start = s.start
            curr_text = text
            curr_end = s.end
        else:
            curr_text += " " + text
            curr_end = s.end
            if (curr_end - curr_start >= 15.0) or len(curr_text) >= 60:
                transcript_blocks.append(f"**[{format_duration(curr_start)} - {format_duration(curr_end)}]**\n\n{curr_text}\n")
                curr_text = ""

    if curr_text:
        transcript_blocks.append(f"**[{format_duration(curr_start)} - {format_duration(curr_end)}]**\n\n{curr_text}\n")

    return transcript_blocks, duration_s


def process_note_ocr(ocr_engine, aweme):
    images = aweme.get('images', [])
    image_texts = []
    for idx, img_obj in enumerate(images):
        url_list = img_obj.get('url_list', [])
        if not url_list:
            continue
        img_url = url_list[0]
        try:
            req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
            img_bytes = urllib.request.urlopen(req, timeout=10).read()
            ocr_res, _ = ocr_engine(img_bytes)
            lines = [line[1].strip() for line in (ocr_res or []) if line[1].strip()]
            if lines:
                block = f"#### 图 {idx+1} 原文识别：\n\n" + "\n\n".join(lines)
                image_texts.append(block)
        except Exception as e:
            image_texts.append(f"#### 图 {idx+1}：(OCR 异常: {e})")
    return image_texts


def collect_items_from_page(dy_page, max_items=500):
    print("[Collector] 开始滚动扫描当前页面的收藏条目...")
    collected_dict = {}
    last_count = 0
    stable_cycles = 0

    for step in range(100):
        items = dy_page.evaluate("""() => {
            const links = Array.from(document.querySelectorAll('a[href*="/video/"], a[href*="/note/"]'));
            return links.map(a => {
                let text = a.innerText.trim();
                let href = a.href;
                let img = a.querySelector('img');
                return {
                    url: href.split('?')[0],
                    raw_href: href,
                    text: text,
                    cover: img ? img.src : ''
                };
            }).filter(x => x.url.includes('/video/') || x.url.includes('/note/'));
        }""")

        for it in items:
            u = it['url']
            if u not in collected_dict:
                collected_dict[u] = it

        current_count = len(collected_dict)
        print(f"  [扫描轮次 {step+1}] 当前已捕获: {current_count} 条", end="\r")

        if current_count >= max_items:
            break

        if current_count == last_count:
            stable_cycles += 1
            if stable_cycles >= 4:
                break
        else:
            stable_cycles = 0

        last_count = current_count
        dy_page.evaluate("window.scrollBy(0, 1500)")
        time.sleep(1.0)

    print()
    res = list(collected_dict.values())
    print(f"[Collector] 扫描完成，共获取 {len(res)} 条有效收藏条目。")
    return res


def export_excel_and_json(detailed_records, output_dir, account_name="抖音收藏"):
    json_path = os.path.join(output_dir, "douyin_favorites.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(detailed_records, f, ensure_ascii=False, indent=2)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "收藏视频清单"
    headers = ["序号", "类型", "视频标题 / 描述", "作者", "发布时间", "时长", "点赞数", "收藏数", "评论数", "标签 Topic", "页面链接"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_font = Font(name="微软雅黑", size=10)
    for r_idx, r in enumerate(detailed_records, start=2):
        ws.append([
            r['index'], r['type'], r['title'], r['author'], r['publish_time'], r['duration'],
            r['likes'], r['collects'], r['comments'], ", ".join(r['tags']), r['video_url']
        ])
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=r_idx, column=col_idx)
            c.font = row_font
            c.border = thin_border
            c.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 5, 6, 7, 8, 9] else "left", vertical="center")

    widths = [8, 10, 45, 18, 18, 10, 12, 12, 12, 25, 35]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    excel_path = os.path.join(output_dir, "douyin_favorites.xlsx")
    wb.save(excel_path)
    return excel_path, json_path


def execute_pipeline(cdp_port=9223, output_dir=None, collected_file=None, whisper_model="small", compute_type="int8"):
    # 自动定位端口
    ports_to_try = [cdp_port] if cdp_port else [9223, 9222]

    with sync_playwright() as p:
        browser = None
        used_port = None
        for pt in ports_to_try:
            try:
                browser = p.chromium.connect_over_cdp(f"http://127.0.0.1:{pt}")
                used_port = pt
                break
            except Exception:
                continue

        if not browser:
            raise ConnectionError(
                f"无法连接到 Chrome 调试端口 (尝试了 {ports_to_try})！\n"
                f"请确保 Chrome 已带 '--remote-debugging-port=9223' 启动且正在运行。"
            )

        context = browser.contexts[0]
        dy_page = next((page for page in context.pages if "douyin.com" in page.url), None)
        if not dy_page:
            raise RuntimeError("未在当前浏览器中检测到已打开的抖音 (douyin.com) 页面，请先在浏览器中登录并进入个人主页！")

        # 尝试嗅探账号身份
        user_info = dy_page.evaluate("""() => {
            let nick = document.querySelector('h1 span')?.innerText || '';
            return { nickname: nick.trim() };
        }""")
        nick_str = sanitize_filename(user_info.get('nickname') or f"account_port_{used_port}")

        if not output_dir:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M")
            output_dir = os.path.join(os.getcwd(), "outputs", f"{nick_str}_{timestamp}")

        os.makedirs(output_dir, exist_ok=True)
        transcripts_dir = os.path.join(output_dir, "transcripts")
        os.makedirs(transcripts_dir, exist_ok=True)

        print(f"==================================================")
        print(f"  Chrome 调试端口 : {used_port}")
        print(f"  检测到账号身份 : {nick_str}")
        print(f"  产物输出目录   : {output_dir}")
        print(f"==================================================")

        # 扫描或载入条目
        if collected_file and os.path.exists(collected_file):
            print(f"[Info] 加载已缓存条目文件: {collected_file}")
            with open(collected_file, "r", encoding="utf-8") as f:
                collected = json.load(f)
        else:
            collected = collect_items_from_page(dy_page)
            # 自动备份条目清单
            cache_file = os.path.join(output_dir, "raw_collected_items.json")
            with open(cache_file, "w", encoding="utf-8") as f:
                json.dump(collected, f, ensure_ascii=False, indent=2)

        if not collected:
            print("[Warning] 未收集到任何有效条目，请确保页面已停留在“我的收藏”标签下。")
            return

        # 阶段 1: 拉取 API 元数据
        print(f"\n[Phase 1] 正在通过会话凭据拉取 {len(collected)} 条完整元数据...")
        detailed_records = []
        raw_aweme_map = {}

        for idx, item in enumerate(collected, start=1):
            url = item['url']
            m = re.search(r'/(video|note)/(\d+)', url)
            if not m:
                continue
            item_type_raw, aweme_id = m.groups()
            item_type = "图文" if item_type_raw == "note" else "视频"

            res = dy_page.evaluate("""async (id) => {
                try {
                    const r = await fetch(`/aweme/v1/web/aweme/detail/?aweme_id=${id}&aid=6383&channel=channel_pc_web&device_platform=webapp`, {
                        credentials: 'include'
                    });
                    return await r.json();
                } catch(e) {
                    return { error: e.toString() };
                }
            }""", aweme_id)

            aweme = res.get('aweme_detail') if isinstance(res, dict) else None
            raw_aweme_map[aweme_id] = aweme

            if aweme:
                desc = aweme.get('desc', '').strip()
                author = aweme.get('author', {})
                stats = aweme.get('statistics', {})
                video = aweme.get('video', {})
                music = aweme.get('music', {})
                tags = re.findall(r'#([^\s#]+)', desc)
                ctime = aweme.get('create_time')
                date_str = datetime.fromtimestamp(ctime).strftime('%Y-%m-%d %H:%M') if ctime else ""
                duration_ms = video.get('duration', 0)
                dur_str = format_duration(duration_ms // 1000)
                play_urls = video.get('play_addr', {}).get('url_list', [])
                download_url = play_urls[0] if play_urls else ""

                rec = {
                    "index": idx,
                    "aweme_id": aweme_id,
                    "type": item_type,
                    "title": desc.replace("\n", " ").strip() if desc else item.get('text', '（无文案）'),
                    "author": author.get('nickname', '未知作者'),
                    "author_sec_uid": author.get('sec_uid', ''),
                    "author_url": f"https://www.douyin.com/user/{author.get('sec_uid', '')}" if author.get('sec_uid') else "",
                    "publish_time": date_str,
                    "duration": dur_str,
                    "likes": stats.get('digg_count', 0),
                    "likes_formatted": format_count(stats.get('digg_count', 0)),
                    "comments": stats.get('comment_count', 0),
                    "comments_formatted": format_count(stats.get('comment_count', 0)),
                    "collects": stats.get('collect_count', 0),
                    "collects_formatted": format_count(stats.get('collect_count', 0)),
                    "shares": stats.get('share_count', 0),
                    "tags": tags,
                    "video_url": url,
                    "download_url": download_url,
                    "music": f"{music.get('title', '')} - {music.get('author', '')}".strip(" -")
                }
            else:
                rec = {
                    "index": idx,
                    "aweme_id": aweme_id,
                    "type": item_type,
                    "title": item.get('text', '（已下架或私密）'),
                    "author": "未知",
                    "author_sec_uid": "",
                    "author_url": "",
                    "publish_time": "",
                    "duration": "00:00",
                    "likes": 0,
                    "likes_formatted": "0",
                    "comments": 0,
                    "comments_formatted": "0",
                    "collects": 0,
                    "collects_formatted": "0",
                    "shares": 0,
                    "tags": re.findall(r'#([^\s#]+)', item.get('text', '')),
                    "video_url": url,
                    "download_url": "",
                    "music": ""
                }

            detailed_records.append(rec)
            if idx % 10 == 0 or idx == len(collected):
                print(f"  [{idx}/{len(collected)}] {rec['author']} | {rec['title'][:30]}")
            time.sleep(0.1)

        # 导出 Excel / JSON
        excel_out, json_out = export_excel_and_json(detailed_records, output_dir, nick_str)
        print(f"[OK] 汇总表格导出成功: {excel_out}")
        print(f"[OK] JSON 数据导出成功: {json_out}")

        # 阶段 2: 逐字稿 ASR 与 OCR
        print(f"\n[Phase 2] 载入 Faster-Whisper ({whisper_model}, {compute_type}) 与 RapidOCR...")
        whisper = WhisperModel(whisper_model, device="cpu", compute_type=compute_type)
        ocr = RapidOCR()

        index_entries = []

        for item in detailed_records:
            idx = item['index']
            author = item['author']
            title = item['title']
            item_type = item['type']
            aweme_id = item['aweme_id']
            video_url = item['video_url']
            publish_time = item['publish_time']
            duration = item.get('duration', '')
            tags = item.get('tags', [])
            stats = f"点赞: {item['likes_formatted']} | 收藏: {item['collects_formatted']} | 评论: {item['comments_formatted']}"

            safe_author = sanitize_filename(author)
            safe_title = sanitize_filename(title[:25])
            filename = f"{idx:02d}_{safe_author}_{safe_title}.md"
            filepath = os.path.join(transcripts_dir, filename)

            index_entries.append({
                "index": idx,
                "type": item_type,
                "author": author,
                "title": title,
                "duration": duration,
                "publish_time": publish_time,
                "filename": filename,
                "tags": tags
            })

            existing_matches = [f for f in os.listdir(transcripts_dir) if f.startswith(f"{idx:02d}_") and f.endswith(".md")]
            if existing_matches:
                match_file = os.path.join(transcripts_dir, existing_matches[0])
                if os.path.getsize(match_file) > 1000:
                    print(f"  [{idx}/{len(detailed_records)}] 已存在跳过: {existing_matches[0]}")
                    continue

            print(f"  [{idx}/{len(detailed_records)}] 正在处理 [{item_type}]: {author} | {title[:28]} (时长: {duration or '图文'})")
            aweme = raw_aweme_map.get(aweme_id)

            md_lines = []
            md_lines.append(f"# {title}\n")
            md_lines.append(f"- **条目序号**：#{idx:02d}")
            md_lines.append(f"- **内容类型**：{item_type}")
            md_lines.append(f"- **创作者**：[{author}]({item.get('author_url', '')})")
            md_lines.append(f"- **发布时间**：{publish_time if publish_time else '未知'}")
            if item_type == "视频" and duration:
                md_lines.append(f"- **视频时长**：{duration}")
            md_lines.append(f"- **互动数据**：{stats}")
            md_lines.append(f"- **原片链接**：[{video_url}]({video_url})")
            if tags:
                md_lines.append(f"- **话题标签**：{' '.join(['#' + t for t in tags])}")
            md_lines.append("\n---\n")

            if not aweme:
                md_lines.append("## 内容状态说明\n")
                md_lines.append("> ⚠️ **该内容当前在抖音已被原作者删除、隐藏或设置为私密，无法拉取原文数据。**\n")
            elif item_type == "图文":
                md_lines.append("## 图文笔记内容整理\n")
                md_lines.append(f"### 笔记描述文案\n\n> {title}\n")
                try:
                    ocr_blocks = process_note_ocr(ocr, aweme)
                    md_lines.append("### 图片逐页 OCR 识别原文\n")
                    md_lines.extend(ocr_blocks)
                except Exception as e:
                    md_lines.append(f"\nOCR 处理异常: {e}\n")
            else:
                play_urls = aweme.get('video', {}).get('play_addr', {}).get('url_list', [])
                fresh_play_url = play_urls[0] if play_urls else ""
                if not fresh_play_url:
                    md_lines.append("## 内容状态说明\n")
                    md_lines.append("> ⚠️ **该视频当前在抖音已被原作者删除、隐藏或设置为私密，无法拉取音轨。**\n")
                else:
                    md_lines.append("## 视频口述逐字原文 (分段带时间戳)\n")
                    try:
                        blocks, dur = process_video_asr(whisper, fresh_play_url)
                        md_lines.extend(blocks)
                    except Exception as e:
                        md_lines.append(f"\n转录异常: {e}\n")

            with open(filepath, "w", encoding="utf-8") as f:
                f.write("\n".join(md_lines))

        # 生成 README.md 目录
        readme_path = os.path.join(transcripts_dir, "README.md")
        readme_lines = [
            f"# 抖音收藏夹视频与图文原文归档目录 ({nick_str})\n",
            f"> 归档时间：{datetime.now().strftime('%Y-%m-%d')} | 总归档条目：{len(index_entries)} 篇\n",
            "| 序号 | 类型 | 标题 | 作者 | 时长/类型 | 发布时间 | 原文文档链接 |",
            "| :---: | :---: | :--- | :--- | :---: | :---: | :--- |"
        ]
        for ent in index_entries:
            tag_str = f" ({', '.join(ent['tags'][:2])})" if ent['tags'] else ""
            dur_str = ent['duration'] if ent['type'] == '视频' else '图文'
            readme_lines.append(f"| {ent['index']} | {ent['type']} | {ent['title'][:35]}{tag_str} | {ent['author']} | {dur_str} | {ent['publish_time'][:10] if ent['publish_time'] else '-'} | [{ent['filename']}](./{ent['filename']}) |")

        with open(readme_path, "w", encoding="utf-8") as f:
            f.write("\n".join(readme_lines))

        print(f"\n[OK] 全流程处理完毕！产物目录: {output_dir}")
        print(f"  - 索引目录: {readme_path}")
        print(f"  - Excel清单: {excel_out}")
        print(f"  - JSON文件: {json_out}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="抖音收藏夹提取与逐字稿一键流水线")
    parser.add_argument("--port", type=int, default=9223, help="Chrome 调试端口 (默认 9223)")
    parser.add_argument("--output", type=str, default=None, help="输出目录 (留空则自动按账号+时间创建)")
    parser.add_argument("--cache", type=str, default=None, help="条目缓存 JSON 文件 (可选)")
    parser.add_argument("--model", type=str, default="small", help="Whisper 模型规格 (默认 small)")
    parser.add_argument("--compute-type", type=str, default="int8", help="量化精度 (默认 int8)")

    args = parser.parse_args()
    execute_pipeline(
        cdp_port=args.port,
        output_dir=args.output,
        collected_file=args.cache,
        whisper_model=args.model,
        compute_type=args.compute_type
    )
